"""European Commission Weekly Oil Bulletin: consumer prices with taxes.

Source page: https://energy.ec.europa.eu/data-and-analysis/weekly-oil-bulletin_en
Reuse: (c) European Union, CC BY 4.0 under Commission Decision 2011/833/EU;
the workbook itself says "Reproduction is authorised provided the source is
acknowledged." Show ATTRIBUTION next to every chart that uses these series.

Two XLSX files are used. Their download URLs carry a fixed document UUID;
the file behind it is replaced each week (the `filename=` query string is
cosmetic and often stale, so it is left off).

- Latest prices with taxes: one sheet, one price date in A2, a row per
  country, columns per product, prices in EUR per 1000 L.
- Price history 2005 onwards: sheet "Prices with taxes", one row per price
  date (newest first), columns named like `DE_price_with_tax_euro95`, also in
  EUR per 1000 L. About 4.5 MB, so it is fetched only for backfill or catch-up.

Prices are as of Monday and published on Thursday. Points are stored at the
price date (UTC midnight), converted to EUR per litre. Blank cells are gaps:
no point is written for them, and never a zero.
"""

import io
import zipfile
from collections.abc import Iterable, Sequence
from datetime import UTC, date, datetime, timedelta

import httpx
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from fuelwatch.http import ProviderError, redact
from fuelwatch.models import Frequency, Point, SeriesSpec

SOURCE = "European Commission, Weekly Oil Bulletin"
ATTRIBUTION = (
    "Source: European Commission, Weekly Oil Bulletin "
    "(energy.ec.europa.eu), © European Union, CC BY 4.0. "
    "Converted from € per 1,000 L to € per litre."
)
PAGE_URL = "https://energy.ec.europa.eu/data-and-analysis/weekly-oil-bulletin_en"
LATEST_URL = "https://energy.ec.europa.eu/document/download/264c2d0f-f161-4ea3-a777-78faae59bea0_en"
HISTORY_URL = (
    "https://energy.ec.europa.eu/document/download/906e60ca-8b6a-44e7-8589-652854d2fd3f_en"
)
HISTORY_SHEET = "Prices with taxes"

BACKFILL_YEARS = 5
# The latest-prices file holds one week. If the caller is further behind
# than this, read the history file instead so no week is skipped.
LATEST_ONLY_WINDOW = timedelta(days=13)

UNIT_IN_FILE = "1000 l"
# Sanity range for a pump price in EUR per litre. Outside it, the file's
# layout or unit has probably changed, so fail loudly rather than store it.
MIN_EUR_PER_L = 0.3
MAX_EUR_PER_L = 6.0

# fuel key -> (history column suffix, header text in the latest file, label)
FUELS = {
    "petrol": ("euro95", "Euro-super 95", "Euro-super 95 petrol"),
    "diesel": ("diesel", "Automotive gas oil", "automotive diesel"),
}
# area key -> (history column prefix, row label match in latest file, name)
AREAS = {
    "eu27": ("EU", "EUR27", "EU average"),
    "de": ("DE", "Germany", "Germany"),
    "fr": ("FR", "France", "France"),
    "it": ("IT", "Italy", "Italy"),
    "es": ("ES", "Spain", "Spain"),
}


def series_id(fuel: str, area: str) -> str:
    return f"eu_pump_{fuel}_{area}"


SERIES: list[SeriesSpec] = [
    SeriesSpec(
        id=series_id(fuel, area),
        name=f"Pump price, {flabel}, {aname} (with taxes)",
        unit="EUR/l",
        source=SOURCE,
        frequency=Frequency.WEEKLY,
        schedule="weekly, Thursday; prices as of the Monday before",
        grace=timedelta(days=2),
    )
    for fuel, (_, _, flabel) in FUELS.items()
    for area, (_, _, aname) in AREAS.items()
]


# ---------------------------------------------------------------- parsing


def read_rows(content: bytes, sheet: str | None = None) -> list[tuple]:
    """Decode an XLSX download into a list of row tuples (cell values)."""
    try:
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except (zipfile.BadZipFile, InvalidFileException, KeyError, OSError, ValueError):
        raise ProviderError("bad_xlsx", "response is not a readable XLSX workbook") from None
    try:
        if sheet is None:
            ws = wb.worksheets[0]
        elif sheet in wb.sheetnames:
            ws = wb[sheet]
        else:
            raise ProviderError("schema", f"sheet {sheet!r} not found; have {wb.sheetnames}")
        return [tuple(r) for r in ws.iter_rows(values_only=True)]
    finally:
        wb.close()


def _utc_midnight(value: object) -> datetime | None:
    if isinstance(value, datetime):
        value = value.date()
    if isinstance(value, date):
        return datetime(value.year, value.month, value.day, tzinfo=UTC)
    return None


def _to_point(sid: str, when: datetime, raw: object) -> Point | None:
    """EUR per 1000 L -> EUR per litre. Blank, non-numeric or non-positive
    cells are gaps and give no point."""
    if isinstance(raw, bool) or not isinstance(raw, int | float) or raw <= 0:
        return None
    eur_per_l = round(raw / 1000.0, 5)
    if not MIN_EUR_PER_L <= eur_per_l <= MAX_EUR_PER_L:
        raise ProviderError(
            "out_of_range", f"{sid} at {when.date()} is {eur_per_l} EUR/l; unit or layout changed?"
        )
    return Point(series=sid, period_end=when, value=eur_per_l)


def parse_latest(rows: Sequence[tuple]) -> list[Point]:
    """Parse the 'latest prices with taxes' sheet."""
    if len(rows) < 3:
        raise ProviderError("schema", "latest-prices sheet has fewer than 3 rows")
    header, units = rows[0], rows[1]
    when = _utc_midnight(units[0] if units else None)
    if when is None:
        raise ProviderError("schema", "latest-prices sheet has no price date in A2")

    cols: dict[str, int] = {}
    for fuel, (_, header_text, _) in FUELS.items():
        idx = [i for i, h in enumerate(header) if isinstance(h, str) and header_text in h]
        if len(idx) != 1:
            raise ProviderError("schema", f"expected one {header_text!r} column, found {len(idx)}")
        if str(units[idx[0]]).strip().lower() != UNIT_IN_FILE:
            raise ProviderError("schema", f"{header_text!r} unit is {units[idx[0]]!r}")
        cols[fuel] = idx[0]

    points: list[Point] = []
    for area, (_, label, _) in AREAS.items():
        matches = [
            r for r in rows[2:] if r and isinstance(r[0], str) and _label_matches(r[0], label)
        ]
        if len(matches) != 1:
            raise ProviderError("schema", f"expected one row for {label!r}, found {len(matches)}")
        row = matches[0]
        for fuel, c in cols.items():
            p = _to_point(series_id(fuel, area), when, row[c] if c < len(row) else None)
            if p:
                points.append(p)
    return points


def _label_matches(cell: str, label: str) -> bool:
    cell = cell.strip()
    if label == "EUR27":
        return "EUR27" in cell
    return cell == label


def parse_history(rows: Sequence[tuple], since: date | None = None) -> list[Point]:
    """Parse the history workbook's 'Prices with taxes' sheet.

    Row 1 holds machine column names, row 3 the units, then one row per
    price date with a datetime in column A. Footnote rows have no date and
    are skipped.
    """
    if len(rows) < 4:
        raise ProviderError("schema", "history sheet has fewer than 4 rows")
    header, units = rows[0], rows[2]

    wanted: dict[int, str] = {}
    for fuel, (suffix, _, _) in FUELS.items():
        for area, (prefix, _, _) in AREAS.items():
            name = f"{prefix}_price_with_tax_{suffix}"
            idx = [i for i, h in enumerate(header) if h == name]
            if len(idx) != 1:
                raise ProviderError("schema", f"expected one column {name!r}, found {len(idx)}")
            if str(units[idx[0]]).strip().lower() != UNIT_IN_FILE:
                raise ProviderError("schema", f"{name!r} unit is {units[idx[0]]!r}")
            wanted[idx[0]] = series_id(fuel, area)

    points: list[Point] = []
    for row in rows[3:]:
        when = _utc_midnight(row[0] if row else None)
        if when is None or (since is not None and when.date() < since):
            continue
        for c, sid in wanted.items():
            p = _to_point(sid, when, row[c] if c < len(row) else None)
            if p:
                points.append(p)
    return points


def latest_date(points: Iterable[Point]) -> datetime | None:
    return max((p.period_end for p in points), default=None)


# ---------------------------------------------------------------- fetching


async def _get_bytes(client: httpx.AsyncClient, url: str) -> bytes:
    try:
        resp = await client.get(url)
        resp.raise_for_status()
    except httpx.HTTPStatusError as e:
        raise ProviderError(
            "http_status", f"{e.response.status_code} from {e.request.url}"
        ) from None
    except httpx.TransportError as e:
        raise ProviderError("transport", f"{type(e).__name__} calling {redact(url)}") from None
    ctype = resp.headers.get("content-type", "")
    if "html" in ctype:
        raise ProviderError("bad_content", f"HTML instead of XLSX from {redact(url)}")
    return resp.content


async def fetch(client: httpx.AsyncClient, since: date | None = None) -> list[Point]:
    """`since=None` backfills five years from the history workbook. A recent
    `since` reads only the small latest-prices file."""
    today = datetime.now(UTC).date()
    if since is not None and today - since <= LATEST_ONLY_WINDOW:
        rows = read_rows(await _get_bytes(client, LATEST_URL))
        points = parse_latest(rows)
        return [p for p in points if p.period_end.date() >= since]

    if since is None:
        since = today - timedelta(days=round(365.25 * BACKFILL_YEARS))
    rows = read_rows(await _get_bytes(client, HISTORY_URL), HISTORY_SHEET)
    return parse_history(rows, since=since)
